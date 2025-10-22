import json
import logging
import re
import datetime
from typing import IO, Union
from io import BytesIO

import psycopg2  # type: ignore
import openpyxl  # type: ignore
from openpyxl.workbook.workbook import Workbook  # type: ignore

import utils  # type: ignore
from models import Lesson, LessonTimings, AbstractScheduleParser  # type: ignore
from parser_app import database

local_timings = LessonTimings.from_json("./parser_app/data/ttsiigh/lesson_timings.json")
institution_id = 1

log = logging.getLogger("parser_app")


class XLSXParser(AbstractScheduleParser):
    def __init__(self, file: IO[bytes] | str):
        super().__init__(file)
        if isinstance(file, bytes):
            self.book: Workbook = openpyxl.load_workbook(filename=BytesIO(file))
        else:
            self.book: Workbook = openpyxl.load_workbook(filename=file)  # type: ignore

    def extract_date(self) -> datetime.date:
        months = {
            "Январь": "01",
            "Февраль": "02",
            "Март": "03",
            "Апрель": "04",
            "Май": "05",
            "Июнь": "06",
            "Июль": "07",
            "Август": "08",
            "Сентябрь": "09",
            "Октябрь": "10",
            "Ноябрь": "11",
            "Декабрь": "12"
        }

        sheet = self.book.worksheets[0]
        date_string = sheet.cell(row=2, column=4).value

        reg = r'(\d{1,2})\s([а-яёА-ЯЁ]+)\s(\d{4})'
        match = re.search(reg, date_string)

        if not match:
            raise Exception(f"Extracting date failed. Date string [{date_string}] does not match expected pattern")

        day = match.group(1)
        month_name = match.group(2)
        year = match.group(3)
        month = months.get(month_name)
        if month:
            formatted_date = f"{day.zfill(2)}.{month}.{year}"
            date_obj = datetime.datetime.strptime(formatted_date, "%d.%m.%Y").date()
            return date_obj
        else:
            raise Exception(f"Invalid month name: {month_name}")

    def __get_group_coord(self) -> utils.Vector2:
        if len(self.book.worksheets) < 1:
            raise Exception("No worksheets found in schedule :(")

        for sheet in self.book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    text = cell.value
                    if text and "Группа" in str(text):
                        return utils.Vector2(cell.row, cell.column)

        raise Exception("It seems like there are no 'Группа' cells in schedule...")

    def extract_groups(self) -> list[str]:
        start_coords = self.__get_group_coord()

        groups = []
        for sheet in self.book.worksheets:
            row = start_coords.x + 1
            while row < sheet.max_row:
                group_cell = sheet.cell(row=row, column=start_coords.y).value
                if group_cell:
                    groups.append(group_cell)
                row += 3

        return groups

    def extract_lessons(self, lesson_timings: LessonTimings) -> list[Lesson]:
        start_coords = self.__get_group_coord()

        ROW_TEACHER_OFFSET = 1
        ROW_CABINET_OFFSET = 2
        ROW_STEP = 3

        sheet = self.book.worksheets[0]
        lessons = []
        row = start_coords.x + 1
        while row < sheet.max_row:
            group_name = sheet.cell(row=row, column=start_coords.y).value or "-"

            for col in range(start_coords.y + 1, sheet.max_column):
                pair_ordinal = (col - 1) - start_coords.y
                pair_name = sheet.cell(row=row, column=col).value or "-"
                pair_teacher = sheet.cell(row=row + ROW_TEACHER_OFFSET, column=col).value or "-"
                pair_cab_num = sheet.cell(row=row + ROW_CABINET_OFFSET, column=col).value or "-"
                pair_cab_num = pair_cab_num if pair_cab_num != "" else "-"

                if pair_name != "-" and pair_name != "Н/Б":
                    lessons.append(Lesson(
                        lesson_timings=lesson_timings,
                        date=self.extract_date(),
                        lesson_num=pair_ordinal,
                        cab_num=str(pair_cab_num).strip(),
                        name=str(pair_name),
                        teacher=str(pair_teacher),
                        institution_id=1,
                        group=group_name
                    ))

            row += ROW_STEP
        return lessons


def get_group_id_by_name(name: str) -> Union[int, None]:
    with psycopg2.connect(
            dsn=utils.DSN
    ) as conn:
        cursor = conn.cursor()
        cursor.execute(
            'SELECT id FROM "groups" WHERE name=%s', (name,)
        )
        group_id = cursor.fetchone()[0]
        return group_id


def add_schedule_from_xlsx(file: IO[bytes]) -> bool:
    parser = XLSXParser(file)
    groups = parser.extract_groups()
    raw_date = parser.extract_date()

    if schedule_exists(raw_date):
        return False

    add_groups(groups)
    lessons = parser.extract_lessons(local_timings)

    database.write_lessons_to_bd(lessons)

    channel = "info_stream:1"
    message = {"type": "new schedule", "date": raw_date.strftime("%Y-%m-%d")}
    utils.redis_client.publish(channel, json.dumps(message))
    return True


def add_groups(groups_list: list[str]):
    with psycopg2.connect(
            dsn=utils.DSN
    ) as conn:
        cursor = conn.cursor()
        for group_name in groups_list:
            cursor.execute('SELECT id FROM "groups" WHERE name = %s', (group_name.lower(),))
            result = cursor.fetchone()
            if not result:
                cursor.execute(
                    'INSERT INTO "groups" (name, institution_id) VALUES (%s, %s)',
                    (group_name.lower(), 1)
                )
                conn.commit()


def schedule_exists(target_date: datetime.date) -> bool:
    query = """
        SELECT 1
        FROM lessons
        WHERE start_time::date <= %s AND end_time::date >= %s AND institution_id = 1 
        LIMIT 1;
    """
    with psycopg2.connect(
            dsn=utils.DSN
    ) as conn:
        cur = conn.cursor()
        cur.execute(query, (target_date, target_date))
        exists = cur.fetchone() is not None
        return exists
